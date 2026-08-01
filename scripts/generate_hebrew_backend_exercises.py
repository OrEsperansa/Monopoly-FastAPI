from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "student-backend-exercises-he"

BLUE = "2563EB"
DARK_BLUE = "1E3A5F"
MUTED = "64748B"
LIGHT_BLUE = "EFF6FF"
LIGHT_GRAY = "F4F6F9"
BORDER = "CBD5E1"
FONT = "Arial"


EXERCISES = [
    {
        "file": "01_getting_started_and_health.docx",
        "title": "תרגיל 1: מתחילים ויוצרים Health Check",
        "level": "בסיסי",
        "goal": "להקים שלד FastAPI תקין, להפעיל שרת מקומי וליצור נקודת קצה ראשונה שניתנת לבדיקה.",
        "result": "בסיום התרגיל השרת עולה, Swagger זמין, והנתיב `GET /api/health` מחזיר תשובה תקינה.",
        "files": ["student-api/app/main.py", "student-api/app/api/game_routes.py", "student-api/app/schemas/game_schemas.py", "student-api/tests/test_api.py"],
        "tasks": [
            "צרו סביבת Python וירטואלית והתקינו את החבילות מתוך `requirements.txt`.",
            "צרו מופע `FastAPI` בתוך `app/main.py` וחברו router עם prefix בשם `/api`.",
            "הגדירו מודל Pydantic בשם `HealthResponse` עם השדה `status`.",
            "ממשו את `GET /api/health` כך שיחזיר בדיוק `status: ok`.",
            "הפעילו CORS כך שממשק React שרץ בפורט אחר יוכל לפנות לשרת באמצעות HTTP.",
            "כתבו בדיקה בשם `test_health` שמוודאת קוד 200 ואת גוף התשובה.",
        ],
        "commands": [
            "cd student-api",
            "python -m venv .venv",
            ".venv\\Scripts\\activate",
            "pip install -r requirements.txt",
            "uvicorn app.main:app --reload",
            "pytest -q",
        ],
        "contract": [
            ("בקשה", "GET /api/health"),
            ("קוד תשובה", "200"),
            ("JSON", '{"status": "ok"}'),
        ],
        "checks": [
            "העמוד `/docs` נפתח ומציג את הנתיב.",
            "`test_health` עובר.",
            "לא נכתב שום חוק משחק ב-route.",
        ],
        "submission": "הגישו את קובצי `app`, את הבדיקה, וצילום מסך של Swagger עם תשובת health.",
    },
    {
        "file": "02_domain_models_and_game_creation.docx",
        "title": "תרגיל 2: מודלים, Repository ויצירת משחק",
        "level": "בסיסי פלוס",
        "goal": "לבנות מנוע משחק טהור ב-Python וליצור משחק חדש עם 2 עד 4 שחקנים ולוח ראשוני.",
        "result": "בסיום התרגיל ניתן ליצור משחק ולשלוף אותו, בלי שהמודלים והשירותים תלויים ב-FastAPI.",
        "files": ["game_engine/models/player.py", "game_engine/models/game.py", "game_engine/models/tile.py", "game_engine/factories/board_factory.py", "game_engine/repositories/game_repository.py", "game_engine/services/game_service.py"],
        "tasks": [
            "צרו `Player` עם `id`, `name`, כסף התחלתי 1500, מיקום 0, דגל פשיטת רגל ורשימת נכסים.",
            "צרו `Game` עם שחקנים, לוח, אינדקס שחקן נוכחי, סטטוס, פעולות זמינות, מנצח ויומן אירועים.",
            "צרו מחלקת בסיס `Tile` עם `id`, `name` ו-`type`.",
            "ממשו `BoardFactory` שמייצר לוח קבוע של 16 משבצות; בשלב זה מותר להשתמש במשבצות בסיסיות.",
            "ממשו `GameRepository` ששומר משחק פעיל יחיד בזיכרון ומאפשר `save`, `get` ו-`reset`.",
            "ממשו `GameService.create_game` ו-`get_game`; השירות מתזמר וה-repository שומר.",
            "חברו `POST /api/game` ו-`GET /api/game` לשירות בלבד.",
        ],
        "commands": ["cd student-api", "pytest -q -k \"create_game or get_game\"", "uvicorn app.main:app --reload"],
        "contract": [
            ("קלט", '{"player_names": ["Ada", "Linus"]}'),
            ("מצב התחלתי", "status=active, current_player_id=1"),
            ("פעולה זמינה", "available_actions=[\"roll\"]"),
            ("גודל לוח", "16"),
        ],
        "checks": [
            "לכל שחקן id רציף שמתחיל ב-1.",
            "המשחק נשמר וניתן לשליפה בבקשת GET נפרדת.",
            "`game_engine` אינו מייבא FastAPI או Pydantic.",
        ],
        "submission": "הגישו את מודלי הדומיין, ה-factory, ה-repository, השירות והבדיקות ליצירה ולשליפה.",
    },
    {
        "file": "03_validation_and_json_contract.docx",
        "title": "תרגיל 3: ולידציה וחוזה JSON מלא",
        "level": "ביניים",
        "goal": "להגדיר חוזה API יציב שמתאים בדיוק לממשק המשותף ולדחות קלט לא תקין לפני הפעלת מנוע המשחק.",
        "result": "בסיום התרגיל כל תשובת משחק מציגה את אותו מבנה, ושמות שחקנים לא תקינים מחזירים HTTP 422.",
        "files": ["app/schemas/game_schemas.py", "app/api/game_routes.py", "game_engine/models/game.py", "tests/test_api.py"],
        "tasks": [
            "הגדירו `CreateGameRequest` עם רשימת `player_names` באורך 2 עד 4.",
            "נקו רווחים בתחילת ובסוף שם ודחו שם ריק או שם שמכיל רווחים בלבד.",
            "דחו שמות כפולים ללא תלות באותיות גדולות וקטנות באנגלית.",
            "הגדירו response models עבור שחקן, משבצת, מצב משחק ותוצאת פעולה.",
            "ממשו `to_dict` במודלי הדומיין בלי להכניס אליהם Pydantic.",
            "הוסיפו `response_model` לכל route כדי ש-FastAPI יאמת גם את הפלט.",
            "כתבו בדיקות לקלט עם שחקן יחיד, שם ריק ושמות כפולים.",
        ],
        "commands": ["cd student-api", "pytest -q -k validation", "pytest -q"],
        "contract": [
            ("GameState", "status, current_player_id, winner_id, available_actions"),
            ("Player", "id, name, money, position, is_bankrupt, owned_tile_ids"),
            ("Tile", "id, name, type, price, rent, owner_id, amount"),
            ("אירועים", "last_event, event_log"),
        ],
        "checks": [
            "קלט לא תקין מחזיר 422 ואינו יוצר משחק.",
            "שדות חסרים במשבצת מיוצגים כ-`null` או אינם נדרשים לפי סכמת התגובה.",
            "`POST /api/game` ו-`GET /api/game` מחזירים אותה צורת GameState.",
        ],
        "submission": "הגישו את סכמות Pydantic, בדיקות הולידציה ודוגמת JSON מלאה של משחק חדש.",
    },
    {
        "file": "04_dice_movement_and_start_bonus.docx",
        "title": "תרגיל 4: קובייה, תנועה ובונוס Start",
        "level": "ביניים",
        "goal": "להוסיף זריקת קובייה, תנועה מעגלית ובונוס מעבר דרך Start באמצעות שירותים קטנים הניתנים לבדיקה.",
        "result": "בסיום התרגיל `POST /api/game/roll` מזיז רק את השחקן הנוכחי ומחזיר מצב משחק מעודכן.",
        "files": ["game_engine/services/dice_service.py", "game_engine/services/movement_service.py", "game_engine/services/game_service.py", "app/api/game_routes.py"],
        "tasks": [
            "ממשו `DiceService.roll` שמחזיר מספר שלם אקראי בין 1 ל-6.",
            "ממשו `MovementService.move_player` שמחשב מיקום חדש באמצעות גודל הלוח.",
            "זהו מעבר דרך Start והוסיפו לשחקן 200; נחיתה רגילה ללא הקפה אינה מעניקה בונוס.",
            "ב-`GameService.roll` ודאו שהפעולה `roll` זמינה לפני שינוי המצב.",
            "עדכנו `last_event` ו-`event_log` עם תוצאת הקובייה והמיקום החדש.",
            "החזירו מ-route הפעולה מעטפת בצורה `state` ולא GameState ישיר.",
            "בבדיקות השתמשו ב-`monkeypatch` כדי שהקובייה תחזיר ערך ידוע.",
        ],
        "commands": ["cd student-api", "pytest -q -k \"roll or movement\""],
        "contract": [
            ("בקשה", "POST /api/game/roll ללא body"),
            ("תשובה", '{"state": { ...GameState... }}'),
            ("דוגמה", "מיקום 14 + תוצאה 3 = מיקום 1 ובונוס 200"),
        ],
        "checks": [
            "בדיקה דטרמיניסטית מאמתת את המיקום לאחר זריקה.",
            "בדיקה נפרדת מאמתת מעבר Start ואת יתרת הכסף.",
            "ה-route אינו מגריל ואינו מחשב מיקום בעצמו.",
        ],
        "submission": "הגישו את שירותי הקובייה והתנועה, פעולת roll ושתי בדיקות דטרמיניסטיות לפחות.",
    },
    {
        "file": "05_tile_inheritance_and_buying.docx",
        "title": "תרגיל 5: ירושה במשבצות וקניית נכסים",
        "level": "ביניים מתקדם",
        "goal": "לייצג סוגי משבצות בירושה ולהוסיף תהליך רכישה חוקי של נכס פנוי.",
        "result": "בסיום התרגיל נחיתה על נכס פנוי מאפשרת buy, ורכישה מעדכנת כסף, בעלות ופעולות זמינות.",
        "files": ["game_engine/models/tile.py", "game_engine/factories/board_factory.py", "game_engine/services/tile_service.py", "game_engine/services/game_service.py"],
        "tasks": [
            "צרו את תתי-המחלקות `StartTile`, `PropertyTile`, `RewardTile`, `PenaltyTile` ו-`SurpriseTile`.",
            "הוסיפו ל-`PropertyTile` מחיר, שכירות ו-`owner_id` שמתחיל כ-`None`.",
            "עדכנו את ה-factory כך שהלוח יכלול 16 משבצות ממספר סוגים ובמזהים 0 עד 15.",
            "בנחיתה על נכס פנוי הגדירו `available_actions` ל-`buy` ו-`end_turn`.",
            "ממשו `GameService.buy`: בדיקת סוג המשבצת, בעלות, מספיק כסף וזמינות הפעולה.",
            "בעת רכישה הורידו מחיר, עדכנו `owner_id`, הוסיפו את id הנכס לשחקן והשאירו רק `end_turn`.",
            "חברו `POST /api/game/buy` וכתבו בדיקה עם קובייה מדומה שמביאה לנכס ידוע.",
        ],
        "commands": ["cd student-api", "pytest -q -k buy"],
        "contract": [
            ("נכס פנוי", "available_actions=[\"buy\", \"end_turn\"]"),
            ("לאחר קנייה", "owner_id=player.id"),
            ("כסף", "money = money - price"),
            ("פעולה הבאה", "available_actions=[\"end_turn\"]"),
        ],
        "checks": [
            "לא ניתן לקנות משבצת שאינה Property.",
            "לא ניתן לקנות נכס תפוס או נכס יקר מדי.",
            "ה-JSON של כל סוג משבצת תואם לחוזה מתרגיל 3.",
        ],
        "submission": "הגישו את היררכיית המשבצות, הלוח המלא, תהליך buy ובדיקות חיוביות ושליליות.",
    },
    {
        "file": "06_tile_effects_rent_and_events.docx",
        "title": "תרגיל 6: שכירות, אפקטים ויומן אירועים",
        "level": "מתקדם",
        "goal": "לרכז את כל תוצאות הנחיתה בשירות משבצות ולהפוך את מצב המשחק למספיק עשיר עבור ממשק תצוגה בלבד.",
        "result": "בסיום התרגיל השרת מטפל בשכירות, פרס, קנס והפתעה ומדווח לממשק מה קרה.",
        "files": ["game_engine/services/tile_service.py", "game_engine/models/game.py", "game_engine/models/player.py", "tests/test_api.py"],
        "tasks": [
            "בנחיתה על נכס של שחקן אחר העבירו את סכום השכירות מהשוכר לבעלים.",
            "בנחיתה על נכס עצמי אל תשנו כסף והשאירו רק את הפעולה `end_turn`.",
            "ב-Reward הוסיפו `amount`; ב-Penalty הורידו `amount`.",
            "ב-Surprise בחרו אפקט באמצעות `random.choice`; כל אפקט חייב להסתיים במצב עקבי.",
            "ממשו `Game.add_event` שמעדכן גם `last_event` וגם מוסיף ל-`event_log`.",
            "הוסיפו הודעה ברורה לכל זריקה, קנייה, תשלום או אפקט. הממשק רק מציג את הטקסט.",
            "בדקו אפקט אקראי באמצעות monkeypatch של `random.choice`.",
        ],
        "commands": ["cd student-api", "pytest -q -k \"rent or reward or penalty or surprise\""],
        "contract": [
            ("שכירות", "השוכר מפסיד rent והבעלים מרוויח rent"),
            ("Reward", "money += amount"),
            ("Penalty", "money -= amount"),
            ("אירוע", "last_event הוא האירוע האחרון והוא מופיע גם בסוף event_log"),
        ],
        "checks": [
            "סכום הכסף הכולל נשמר בעת תשלום שכירות בין שחקנים.",
            "כל נחיתה משאירה רשימת פעולות חוקית.",
            "אין ב-React חישוב שכירות, בונוס, קנס או הפתעה.",
        ],
        "submission": "הגישו את TileService, מנגנון האירועים ובדיקות לכל ארבעת סוגי האפקטים.",
    },
    {
        "file": "07_turns_actions_and_errors.docx",
        "title": "תרגיל 7: תורות, פעולות זמינות ושגיאות",
        "level": "מתקדם",
        "goal": "להגן על מכונת המצבים של המשחק כך שניתן לבצע רק פעולה חוקית ובזמן הנכון.",
        "result": "בסיום התרגיל השרת מנהל תורות, מדלג על שחקנים לא פעילים ומחזיר שגיאות HTTP עקביות.",
        "files": ["game_engine/services/turn_service.py", "game_engine/exceptions.py", "app/exception_handlers.py", "game_engine/services/game_service.py"],
        "tasks": [
            "ממשו `TurnService.advance_turn` שמתקדם בצורה מעגלית לשחקן הפעיל הבא.",
            "ממשו `POST /api/game/end-turn` והחזירו לאחריו `available_actions=[\"roll\"]`.",
            "לפני roll, buy או end_turn בדקו שהפעולה נמצאת ב-`available_actions`.",
            "צרו `GameError` בסיסי, `GameNotFoundError` עם status 404 ו-`InvalidActionError` עם status 409.",
            "רשמו exception handler ב-FastAPI שמתרגם חריגת דומיין ל-JSON בצורה `detail`.",
            "ודאו שפעולה שנכשלה אינה משנה כסף, מיקום, בעלות או תור.",
            "כתבו בדיקות לסבב תורות, לדילוג על bankrupt ולפעולה לא חוקית.",
        ],
        "commands": ["cd student-api", "pytest -q -k \"turn or action or error\""],
        "contract": [
            ("אין משחק", '404 + {"detail": "..."}'),
            ("פעולה לא זמינה", '409 + {"detail": "..."}'),
            ("סיום תור", "current_player_id מתחלף והפעולה הבאה היא roll"),
        ],
        "checks": [
            "לא ניתן לזרוק פעמיים באותו תור.",
            "לא ניתן לקנות כאשר buy אינו זמין.",
            "שחקן bankrupt אינו מקבל תור.",
            "`game_engine` עדיין אינו מכיר קודי HTTP.",
        ],
        "submission": "הגישו את TurnService, החריגות, ה-handler ובדיקות שמוכיחות ששגיאה אינה משנה מצב.",
    },
    {
        "file": "08_complete_game_and_ui_integration.docx",
        "title": "תרגיל 8: משחק מלא וחיבור לממשק",
        "level": "מסכם",
        "goal": "להשלים פשיטת רגל, מנצח ואיפוס, ולאמת שהשרת מפעיל משחק מלא מול ממשק React פסיבי.",
        "result": "בסיום התרגיל כל המשחק עובד דרך ה-API; הממשק רק שולח בקשות ומציג את מצב השרת.",
        "files": ["game_engine/services/turn_service.py", "game_engine/services/game_service.py", "app/api/game_routes.py", "app/main.py", "tests/test_api.py"],
        "tasks": [
            "כאשר כסף של שחקן יורד מתחת ל-0, סמנו `is_bankrupt=True`.",
            "לאחר כל פעולה שמשנה כסף בדקו כמה שחקנים פעילים נשארו.",
            "כאשר נשאר שחקן פעיל אחד, הגדירו `status=finished`, שמרו `winner_id` ורוקנו `available_actions`.",
            "ממשו `POST /api/game/reset` שמחזיר משחק במצב `not_started` עם players ו-board ריקים.",
            "הריצו את כל הבדיקות ותקנו הפרות חוזה בלי לשנות את טיפוסי TypeScript בממשק.",
            "הריצו backend ו-`shared-ui`, הזינו בממשק את כתובת ה-API ושחקו לפחות חמישה תורות.",
            "בדקו שכל כפתור מוצג לפי `available_actions`, שכל הטקסט מגיע מ-`last_event`/`event_log`, ושאין חישוב חוקי משחק ב-React.",
            "בצעו משחק בדיקה עד מנצח ולאחר מכן reset ויצירת משחק חדש.",
        ],
        "commands": [
            "cd student-api",
            "pytest -q",
            "uvicorn app.main:app --reload",
            "cd ../shared-ui",
            "npm install",
            "npm run dev",
            "npm run build",
        ],
        "contract": [
            ("GET", "/api/health, /api/game"),
            ("POST", "/api/game, /api/game/roll, /api/game/buy"),
            ("POST נוסף", "/api/game/end-turn, /api/game/reset"),
            ("מצב סופי", "status=finished, winner_id=<id>, available_actions=[]"),
            ("איפוס", "status=not_started, players=[], board=[]"),
        ],
        "checks": [
            "כל בדיקות `pytest` עוברות.",
            "`npm run build` עובר ללא שינוי חוזה הממשק.",
            "ניתן להשלים רצף create, roll, buy/skip, end-turn ו-reset דרך ה-UI.",
            "השרת הוא מקור האמת היחיד לכסף, מיקום, בעלות, תור, פעולות ומנצח.",
            "Routes עוסקים רק ב-HTTP; חוקי המשחק נמצאים בשירותי `game_engine`.",
        ],
        "submission": "הגישו backend מלא, תוצאות `pytest` ו-`npm run build`, וצילומי מסך של משחק פעיל, event log ומסך ניצחון.",
    },
]


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_rtl(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def format_run(run, *, code=False, bold=False, size=None, color=None):
    run.font.name = "Consolas" if code else FONT
    run.font.size = Pt(size or 11)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    face = "Consolas" if code else FONT
    for attr in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{attr}"), face)
    if re.search(r"[\u0590-\u05FF]", run.text):
        if r_pr.find(qn("w:rtl")) is None:
            r_pr.append(OxmlElement("w:rtl"))
        lang = r_pr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            r_pr.append(lang)
        lang.set(qn("w:bidi"), "he-IL")


def add_rich_text(paragraph, text, *, bold=False, size=None, color=None):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        run = paragraph.add_run(part[1:-1] if is_code else part)
        format_run(run, code=is_code, bold=bold, size=size, color=color)


def add_para(doc, text, *, style=None, bold=False, size=None, color=None, after=6, keep=False):
    paragraph = doc.add_paragraph(style=style)
    set_rtl(paragraph)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.keep_with_next = keep
    add_rich_text(paragraph, text, bold=bold, size=size, color=color)
    return paragraph


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_rtl(paragraph)
    add_rich_text(paragraph, text, bold=True)
    return paragraph


def add_list(doc, items, reference):
    num_id = "11" if reference == "numbers" else "10"
    for item in items:
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), num_id)
        num_pr.append(ilvl)
        num_pr.append(num_id_el)
        p_pr.insert(0, num_pr)
        add_rich_text(p, item)


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    for abstract_id, num_id, fmt, marker in ((10, 10, "bullet", "•"), (11, 11, "decimal", "%1.")):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), marker)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "right")
        p_pr = OxmlElement("w:pPr")
        bidi = OxmlElement("w:bidi")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "260")
        tabs.append(tab)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:right"), "540")
        indent.set(qn("w:hanging"), "280")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([bidi, tabs, spacing, indent])
        lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)


def add_code_block(doc, lines):
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)
        p.paragraph_format.space_before = Pt(0 if index else 5)
        p.paragraph_format.space_after = Pt(0 if index < len(lines) - 1 else 7)
        p_pr = p._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), LIGHT_GRAY)
        p_pr.append(shading)
        run = p.add_run(line)
        format_run(run, code=True, size=9.5)


def add_contract_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_widths = [Inches(1.75), Inches(4.75)]
    table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    tbl_w = table._tbl.tblPr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    table._tbl.tblPr.append(tbl_ind)
    headers = ["רכיב", "דרישה"]
    for cell, text, width in zip(table.rows[0].cells, headers, table_widths):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EEF5")
        cell._tc.get_or_add_tcPr().append(shading)
        p = cell.paragraphs[0]
        set_rtl(p)
        add_rich_text(p, text, bold=True)
    for label, value in rows:
        cells = table.add_row().cells
        for cell, text, width in zip(cells, (label, value), table_widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            set_rtl(p)
            add_rich_text(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    set_rtl(header)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_rich_text(header, "Property Empire | פרויקט Backend ב-FastAPI", size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("עמוד ")
    format_run(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    add_numbering(doc)


def add_title_block(doc, exercise):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    add_rich_text(p, f"תרגיל {EXERCISES.index(exercise) + 1} מתוך {len(EXERCISES)}", bold=True, size=10, color=BLUE)
    title = doc.add_paragraph()
    set_rtl(title)
    title.paragraph_format.space_after = Pt(5)
    add_rich_text(title, exercise["title"], bold=True, size=23, color=DARK_BLUE)
    subtitle = doc.add_paragraph()
    set_rtl(subtitle)
    subtitle.paragraph_format.space_after = Pt(16)
    add_rich_text(subtitle, f"רמת קושי: {exercise['level']} | תרגיל מצטבר", size=11, color=MUTED)

    callout = doc.add_paragraph()
    set_rtl(callout)
    callout.paragraph_format.left_indent = Inches(0.12)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_after = Pt(12)
    p_pr = callout._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shading)
    add_rich_text(callout, "כלל ארכיטקטוני: ממשק React אינו מחשב חוקי משחק. הוא שולח בקשות HTTP ומציג את ה-JSON שהשרת מחזיר.", bold=True, color=DARK_BLUE)


def build_exercise(exercise):
    doc = Document()
    style_document(doc)
    add_title_block(doc, exercise)

    add_heading(doc, "מטרת התרגיל", 1)
    add_para(doc, exercise["goal"])
    add_para(doc, "תוצר צפוי: " + exercise["result"], bold=True)

    add_heading(doc, "לפני שמתחילים", 1)
    add_list(doc, [
        "השלימו קודם את כל התרגילים הקודמים; כל שלב מסתמך על החוזה שכבר נבנה.",
        "עבדו רק בתוך `student-api` אלא אם סעיף מפורש מבקש בדיקת אינטגרציה מול `shared-ui`.",
        "אל תשנו את ממשק React ואת טיפוסי הנתונים שלו כדי לעקוף כשל בשרת.",
        "הריצו בדיקות לפני השינוי ואחריו, ושמרו על בדיקות קיימות שעוברות.",
    ], "bullets")

    add_heading(doc, "קבצים מרכזיים", 1)
    add_list(doc, exercise["files"], "bullets")

    add_heading(doc, "משימות לביצוע", 1)
    add_list(doc, exercise["tasks"], "numbers")

    add_heading(doc, "חוזה והתנהגות נדרשת", 1)
    add_contract_table(doc, exercise["contract"])

    add_heading(doc, "פקודות עבודה", 1)
    add_code_block(doc, exercise["commands"])

    add_heading(doc, "בדיקות קבלה", 1)
    add_list(doc, exercise["checks"], "bullets")

    add_heading(doc, "הגשה", 1)
    add_para(doc, exercise["submission"])
    add_para(doc, "לפני ההגשה: ודאו שכל הבדיקות שנוספו עד שלב זה עוברות ושאין חוק משחק בתוך `app/api/game_routes.py` או בתוך `shared-ui`.", bold=True, color=DARK_BLUE)

    core = doc.core_properties
    core.title = exercise["title"]
    core.subject = "Property Empire - תרגילי FastAPI לתלמידים"
    core.author = "Property Empire Classroom"
    core.keywords = "FastAPI, Python, API, משחק, תרגיל"

    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / exercise["file"]
    doc.save(target)
    return target


def main():
    targets = [build_exercise(exercise) for exercise in EXERCISES]
    for target in targets:
        print(target)


if __name__ == "__main__":
    main()
