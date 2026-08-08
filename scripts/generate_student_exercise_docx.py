from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = ROOT / "student-tests"
OUTPUT_ROOT = ROOT / "docs" / "student-backend-exercises-he"

FONT = "Arial"
MONO_FONT = "Consolas"
FOREST = "143D36"
FOREST_DARK = "0F2926"
GOLD = "D6A84F"
INK = "17201D"
PAPER = "F4EAD5"
PAPER_LIGHT = "FBF7EF"
MUTED = "5D6B66"
LINE = "D8C9AE"
WHITE = "FFFFFF"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_START_END = 120


@dataclass(frozen=True)
class Block:
    kind: str
    text: str = ""
    rows: tuple[tuple[str, ...], ...] = ()


EXERCISE_FILES = {
    1: "01-health-check.docx",
    2: "02-create-game.docx",
    3: "03-retrieve-game-state.docx",
    4: "04-roll-movement-and-start-bonus.docx",
    5: "05-end-turn.docx",
    6: "06-tile-effects.docx",
    7: "07-buy-property.docx",
    8: "08-rent.docx",
    9: "09-complete-game.docx",
}


SETUP_BLOCKS = [
    Block("title", "Sigit Empire Academy"),
    Block("subtitle", "מדריך הקמה, עבודה והרצת תרגילי צד השרת"),
    Block(
        "lead",
        "במסלול הזה בונים בהדרגה שרת FastAPI למשחק לוח. כל תרגיל מוסיף יכולת אחת, "
        "והבדיקות מתייחסות לשרת כקופסה שחורה דרך HTTP בלבד.",
    ),
    Block("h1", "איך סביבת העבודה בנויה"),
    Block(
        "paragraph",
        "המערכת מורכבת משלושה חלקים עצמאיים. ה-Backend הוא הקוד שאתם כותבים; ה-UI מציג את מצב המשחק ושולח פעולות; "
        "ו-Test Runner מריץ את בדיקות הקבלה של התרגיל שבחרתם.",
    ),
    Block("bullet", "Student API — שרת FastAPI שאתם מממשים ומפעילים, בדרך כלל ב-http://localhost:8000."),
    Block("bullet", "Game UI — ממשק המשחק והמעבדה, זמין כברירת מחדל ב-http://localhost:5187."),
    Block("bullet", "Test Runner — שירות שמריץ את בדיקות ה-HTTP, זמין כברירת מחדל ב-http://localhost:9001."),
    Block("note", "מקור האמת הוא תמיד מצב המשחק שה-Backend מחזיר. ה-UI אינו מחשב תנועה, כסף, בעלות, שכירות, תורות או ניצחון."),
    Block("h1", "מה צריך לפני שמתחילים"),
    Block("bullet", "Python 3.13 או גרסה תואמת, כולל pip ויכולת ליצור virtual environment."),
    Block("bullet", "Docker Desktop לצורך הפעלת ה-UI וה-Test Runner באמצעות containers."),
    Block("bullet", "עותק מקומי של מאגר הפרויקט."),
    Block("h1", "הפעלת ה-Backend"),
    Block("number", "פתחו terminal בתיקיית student-api."),
    Block("code", "python -m venv .venv\n.venv\\Scripts\\activate\npip install -r requirements.txt"),
    Block("number", "צרו ASGI application במבנה שתבחרו. התרגילים אינם מחייבים שמות תיקיות, מחלקות או Design Pattern מסוים, אלא אם התרגיל עצמו אומר אחרת."),
    Block("number", "הפעילו את השרת על פורט זמין. לדוגמה:"),
    Block("code", "uvicorn app.main:app --reload --port 8000"),
    Block("number", "ודאו שהכתובת http://localhost:8000/api/health מחזירה את התשובה שנדרשת בתרגיל 1."),
    Block("h1", "הפעלת הממשק ומעבדת הבדיקות"),
    Block("number", "ודאו ש-Docker Desktop פועל."),
    Block("number", "מתיקיית השורש של המאגר משכו והפעילו את ה-images:"),
    Block("code", "docker compose pull\ndocker compose up -d"),
    Block("number", "פתחו בדפדפן את http://localhost:5187."),
    Block("number", "הזינו את כתובת ה-API שלכם, בדרך כלל http://localhost:8000, ולחצו על התחברות."),
    Block("number", "במסך המשחק אפשר ליצור משחק ולשלוח רק פעולות שהשרת החזיר בתוך available_actions."),
    Block("number", "במסך הבדיקות בחרו תרגיל. מצב הרצה מצטברת מריץ גם את כל התרגילים שקדמו לו."),
    Block("h1", "הרצת בדיקות משורת הפקודה"),
    Block("paragraph", "אפשר להריץ תרגיל ואת כל דרישות הקדם שלו גם ללא ה-UI, מתיקיית השורש של המאגר:"),
    Block("code", "python student-tests/run_checkpoint.py 4 --api-url http://localhost:8000"),
    Block(
        "paragraph",
        "בכישלון מוצגים הפונקציונליות שנבדקה, נקודת הקצה, הקלט, מיקום הערך בתוך ה-JSON, הערך הצפוי והערך שהתקבל.",
    ),
    Block("h1", "מסלול התרגילים"),
    Block(
        "table",
        rows=(
            ("תרגיל", "נושא", "תוצאה מרכזית"),
            ("1", "בדיקת תקינות", "שירות פעיל שמחזיר JSON"),
            ("2", "יצירת משחק", "GameState התחלתי מלא"),
            ("3", "שמירת מצב", "שליפת המשחק הפעיל"),
            ("4", "קובייה ותנועה", "תנועה מעגלית ובונוס Start"),
            ("5", "סיום תור", "מעבר לשחקן הבא"),
            ("6", "השפעות משבצות", "Reward, Penalty, Surprise ו-Property"),
            ("7", "רכישת נכס", "עדכון כסף ובעלות"),
            ("8", "שכירות", "העברת כסף בין שחקנים"),
            ("9", "משחק מלא", "פשיטת רגל, מנצח ואיפוס"),
        ),
    ),
    Block("h1", "שיטת עבודה מומלצת"),
    Block("number", "קראו את כל ה-README של התרגיל לפני כתיבת קוד."),
    Block("number", "ממשו רק את היכולת החדשה ושמרו על החוזים מהתרגילים הקודמים."),
    Block("number", "הפעילו את ה-Backend והריצו את התרגיל הנוכחי."),
    Block("number", "כאשר בדיקה נכשלת, השוו את מיקום ה-JSON והערכים להוראה המתאימה במסמך."),
    Block("number", "לפני מעבר לתרגיל הבא, הריצו בדיקה מצטברת כדי לוודא שלא נשברה התנהגות קודמת."),
    Block("h1", "פתרון תקלות נפוצות"),
    Block("h2", "ה-UI אינו מצליח להתחבר ל-API"),
    Block("bullet", "ודאו שהשרת פועל ושכתובת ה-API בממשק כוללת את הפורט הנכון."),
    Block("bullet", "אם הדפדפן מדווח על CORS, הוסיפו ל-FastAPI את CORSMiddleware ואפשרו את כתובת ה-UI."),
    Block("h2", "הפורט כבר תפוס"),
    Block("bullet", "הפעילו את השירות בפורט אחר ועדכנו בהתאם את כתובת ה-API בממשק או את משתני הפורטים של Docker Compose."),
    Block("h2", "בדיקות רבות נשלחות ברצף"),
    Block("bullet", "תרגילי המשחק המלא מריצים רצף תורות אמיתי. אם מתקבלת שגיאה, קראו את הכישלון הראשון לפני הרצה חוזרת."),
    Block("h2", "שינוי קוד אינו מופיע"),
    Block("bullet", "הפעילו מחדש את שרת ה-Backend. לאחר שינוי ב-images, בנו או משכו אותם מחדש והפעילו docker compose up -d."),
]


def strip_inline_markdown(text: str) -> str:
    text = re.sub(r"!\[([^]]*)]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)]\(([^)]+)\)", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"(?<!`)`([^`]+)`(?!`)", r"\1", text)
    return text.strip()


def split_table_row(line: str) -> tuple[str, ...]:
    stripped = line.strip().strip("|")
    return tuple(strip_inline_markdown(cell.strip()) for cell in stripped.split("|"))


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def parse_markdown(markdown: str) -> list[Block]:
    lines = markdown.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            index += 1
            blocks.append(Block("code", "\n".join(code_lines)))
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            blocks.append(Block("table", rows=tuple(rows)))
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            blocks.append(Block("title" if level == 1 else f"h{level - 1}", heading.group(2).strip()))
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            blocks.append(Block("bullet", bullet.group(1).strip()))
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            blocks.append(Block("number", numbered.group(1).strip()))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                candidate.startswith("```")
                or candidate.startswith("#")
                or re.match(r"^-\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
                or (candidate.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]))
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        blocks.append(Block("paragraph", " ".join(paragraph_lines).strip()))

    return blocks


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", CELL_MARGIN_TOP_BOTTOM),
        ("bottom", CELL_MARGIN_TOP_BOTTOM),
        ("start", CELL_MARGIN_START_END),
        ("end", CELL_MARGIN_START_END),
    ):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, *, font: str = FONT, size: float = 11, color: str = INK, bold: bool | None = None, italic: bool = False) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.italic = italic
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attribute}"), font)
    language = r_pr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        r_pr.append(language)
    language.set(qn("w:bidi"), "he-IL")


def set_paragraph_rtl(paragraph, *, alignment=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.insert(0, OxmlElement("w:bidi"))


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def shade_paragraph(paragraph, fill: str, border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        right = OxmlElement("w:right")
        right.set(qn("w:val"), "single")
        right.set(qn("w:sz"), "18")
        right.set(qn("w:space"), "8")
        right.set(qn("w:color"), border)
        p_bdr.append(right)
        p_pr.append(p_bdr)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.extend([begin, instruction, separate, display, end])
    paragraph._p.append(field_run)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)

    for style_name, size, before, after, color in (
        ("Title", 28, 0, 8, FOREST),
        ("Subtitle", 13, 0, 18, MUTED),
        ("Heading 1", 16, 18, 10, FOREST),
        ("Heading 2", 13, 14, 7, FOREST),
        ("Heading 3", 12, 10, 5, FOREST_DARK),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)


def add_numbering(document: Document, *, decimal: bool) -> int:
    numbering = document.part.numbering_part.element
    existing_nums = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_nums, default=0) + 1
    existing_abstract = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing_abstract, default=-1) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if decimal else "bullet")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if decimal else "\uf0b7")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "right")
    p_pr = OxmlElement("w:pPr")
    p_pr.append(OxmlElement("w:bidi"))
    indent = OxmlElement("w:ind")
    indent.set(qn("w:right"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.extend([start, num_fmt, level_text, level_jc, p_pr])
    if not decimal:
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Symbol")
        r_fonts.set(qn("w:hAnsi"), "Symbol")
        r_pr.append(r_fonts)
        level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.extend([level, number])
    p_pr.append(num_pr)


def configure_document(document: Document, running_label: str) -> tuple[int, int]:
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    sect_pr = section._sectPr
    if sect_pr.find(qn("w:bidi")) is None:
        sect_pr.append(OxmlElement("w:bidi"))

    header = section.header
    header.is_linked_to_previous = False
    header_paragraph = header.paragraphs[0]
    set_paragraph_rtl(header_paragraph)
    header_paragraph.paragraph_format.space_after = Pt(0)
    brand = header_paragraph.add_run("SIGIT EMPIRE ACADEMY")
    set_run_font(brand, size=9, color=GOLD, bold=True)
    label = header_paragraph.add_run(f"  ·  {running_label}")
    set_run_font(label, size=9, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.paragraph_format.space_before = Pt(0)
    add_page_number(footer_paragraph)

    return 0, 0


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def add_inline_runs(paragraph, text: str, *, base_size: float = 11, base_color: str = INK, bold: bool = False) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=base_size, color=base_color, bold=bold)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, font=MONO_FONT, size=base_size - 0.5, color=FOREST_DARK, bold=False)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EFE9DD")
            run._element.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, color=base_color, bold=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size, color=base_color, bold=bold)


def add_code_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.15
    shade_paragraph(paragraph, "EEF2F0", FOREST)
    lines = text.split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, font=MONO_FONT, size=9.5, color=FOREST_DARK)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_width.set(qn("w:type"), "dxa")
    tbl_indent = tbl_pr.find(qn("w:tblInd"))
    if tbl_indent is None:
        tbl_indent = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_indent)
    tbl_indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_indent.set(qn("w:type"), "dxa")
    if tbl_pr.find(qn("w:bidiVisual")) is None:
        tbl_pr.append(OxmlElement("w:bidiVisual"))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def table_widths(column_count: int) -> list[int]:
    if column_count == 4:
        return [700, 2800, 1600, 4260]
    if column_count == 3:
        return [1100, 3000, 5260]
    if column_count == 2:
        return [1800, 7560]
    base = CONTENT_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_table(document: Document, rows: tuple[tuple[str, ...], ...]) -> None:
    if not rows:
        return
    column_count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=column_count)
    set_table_geometry(table, table_widths(column_count))
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_fill(cell, FOREST if row_index == 0 else (PAPER_LIGHT if row_index % 2 else WHITE))
            paragraph = cell.paragraphs[0]
            set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(value)
            set_run_font(run, size=9.5, color=WHITE if row_index == 0 else INK, bold=row_index == 0)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_block(document: Document, block: Block, bullet_num_id: int, decimal_num_id: int) -> None:
    if block.kind == "title":
        paragraph = document.add_paragraph(style="Title")
        set_paragraph_rtl(paragraph)
        add_inline_runs(paragraph, block.text, base_size=28, base_color=FOREST, bold=True)
        return
    if block.kind == "subtitle":
        paragraph = document.add_paragraph(style="Subtitle")
        set_paragraph_rtl(paragraph)
        add_inline_runs(paragraph, block.text, base_size=13, base_color=MUTED)
        return
    if block.kind in {"h1", "h2", "h3"}:
        paragraph = document.add_paragraph(style={"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[block.kind])
        set_paragraph_rtl(paragraph)
        size = {"h1": 16, "h2": 13, "h3": 12}[block.kind]
        add_inline_runs(paragraph, block.text, base_size=size, base_color=FOREST if block.kind != "h3" else FOREST_DARK, bold=True)
        return
    if block.kind == "code":
        add_code_block(document, block.text)
        return
    if block.kind == "table":
        add_table(document, block.rows)
        return

    list_style = {"bullet": "List Bullet", "number": "List Number"}.get(block.kind)
    paragraph = document.add_paragraph(style=list_style)
    set_paragraph_rtl(paragraph)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if block.kind in {"lead", "note"}:
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(10)
        shade_paragraph(paragraph, PAPER if block.kind == "lead" else "EEF2F0", GOLD if block.kind == "lead" else FOREST)
    add_inline_runs(paragraph, block.text, bold=block.kind == "lead")


def build_document(blocks: Iterable[Block], target: Path, running_label: str) -> None:
    document = Document()
    bullet_num_id, decimal_num_id = configure_document(document, running_label)
    for block in blocks:
        add_block(document, block, bullet_num_id, decimal_num_id)
    document.core_properties.title = running_label
    document.core_properties.subject = "Sigit Empire Academy — Backend exercises"
    document.core_properties.author = "Sigit Empire Academy"
    document.core_properties.keywords = "FastAPI, Backend, HTTP, Sigit Empire"
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def visible_source_content(blocks: Iterable[Block]) -> Counter[str]:
    values: list[str] = []
    for block in blocks:
        if block.kind == "table":
            values.extend(cell for row in block.rows for cell in row if cell)
        elif block.text:
            values.append(strip_inline_markdown(block.text))
    return Counter(value for value in values if value)


def visible_docx_content(path: Path) -> Counter[str]:
    document = Document(path)
    values = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    values.append(value)
    return Counter(values)


def verify_exact_content(source_blocks: list[Block], target: Path) -> None:
    expected = visible_source_content(source_blocks)
    actual = visible_docx_content(target)
    if expected != actual:
        missing = expected - actual
        extra = actual - expected
        raise AssertionError(f"DOCX content differs from source for {target.name}; missing={missing}; extra={extra}")


def generate_exercises() -> list[Path]:
    generated: list[Path] = []
    for number, file_name in EXERCISE_FILES.items():
        source = SOURCE_ROOT / f"exercise_{number:02d}" / "README.md"
        blocks = parse_markdown(source.read_text(encoding="utf-8"))
        target = OUTPUT_ROOT / file_name
        build_document(blocks, target, f"תרגיל {number}")
        verify_exact_content(blocks, target)
        generated.append(target)
    return generated


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    expected_names = {"00-setup.docx", *EXERCISE_FILES.values()}
    for existing in OUTPUT_ROOT.glob("*.docx"):
        if existing.name not in expected_names:
            existing.unlink()

    setup_target = OUTPUT_ROOT / "00-setup.docx"
    build_document(SETUP_BLOCKS, setup_target, "מדריך התחלה")
    generated = [setup_target, *generate_exercises()]
    print("\n".join(str(path.relative_to(ROOT)) for path in generated))


if __name__ == "__main__":
    main()
